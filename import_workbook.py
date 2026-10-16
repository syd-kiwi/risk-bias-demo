from pathlib import Path
import openpyxl, json, shutil, sys
src=Path(sys.argv[1]); out=Path(__file__).parent
w=openpyxl.load_workbook(src,data_only=True); s=w['Sheet1']; records=[]
def val(c): return str(c.value).strip() if c.value is not None else ''
for row in s.iter_rows(min_row=2):
 if row[0].value:
  rec=dict(name=val(row[0]),description=val(row[1]),citation=val(row[2]),url=row[3].hyperlink.target if row[3].hyperlink else val(row[3]),studyTypes=[],evidence=val(row[5]),notes=val(row[6]),sourceRow=row[0].row)
  records.append(rec)
 if records and row[4].value:
  typ=val(row[4])
  if typ not in records[-1]['studyTypes']: records[-1]['studyTypes'].append(typ)
assert len(records)==sum(bool(s.cell(i,1).value) for i in range(2,s.max_row+1))
(out/'data.js').write_text('window.TOOL_DATA = '+json.dumps(records,ensure_ascii=False,indent=2)+';\n')
shutil.copyfile(src,out/'source-workbook.xlsx')
print(json.dumps({'records':len(records),'studyTypes':len(set(t for r in records for t in r['studyTypes'])),'evidence':sorted(set(r['evidence'] for r in records))}))
